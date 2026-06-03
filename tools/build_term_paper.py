from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = r"D:\GitRepository\OSP2\OSP-Archive\docs\期末论文_多人RDW主动重置仲裁方法.docx"
RESULT_DIR = r"D:\zsm\Desktop\OSP数据处理\OSP_APF_10_P4\multigroup_user_scaling_figures"


def set_run(run, font="Times New Roman", size=10.5, bold=False, italic=False):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(doc, text="", align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run(r, size=14 if level == 1 else 12, bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run(r)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="BFBFBF", sz="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, "E8EEF5")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run(r, size=9.5, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_width(cell, widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run(r, size=9.5)
    return table


def add_figure(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(15.2))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    set_run(r, size=9.5, italic=True)


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)

    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 10.5)):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        st.paragraph_format.space_after = Pt(6 if name == "Heading 1" else 4)


def add_front(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("面向局部冲突恢复的多人重定向行走主动重置仲裁方法")
    set_run(r, size=18, bold=True)

    add_heading(doc, "摘要", 1)
    add_para(doc, "多人重定向行走需要在有限物理空间中支持多个用户自然行走。现有用户间重置机制通常在两名用户接近安全阈值时被动触发，能够避免真实碰撞，但不直接评价冲突处理后的恢复质量。本文围绕这一问题设计一种面向局部冲突恢复的主动重置仲裁方法。该方法把双边用户重置、短距离连续重置和低恢复行走距离视为高代价恢复事件；在用户 pair 仍有正恢复余量、但恢复余量持续下降时进入主动干预窗口；随后用局部 max-min 仲裁选择单侧用户执行重置。本文不设计新的重置方向策略，而沿用原方向策略，以单独检验触发时机和用户选择的作用。基于 3、4、5、6 用户 APF_OSP 仿真实验结果，Active Reset 没有增加总 reset rows，并在所有用户数下提高 D_recover 中位数；在 4、5、6 用户下，短距离恢复风险显著下降。结果说明，主动重置可以替代部分高代价被动恢复，而不是简单增加额外打断。")
    add_para(doc, "关键词：重定向行走；多人虚拟现实；主动重置；局部冲突恢复；APF_OSP；恢复距离")

    add_heading(doc, "Abstract", 1)
    add_para(doc, "Multi-user redirected walking allows multiple users to walk naturally within a limited physical space. Existing user-user reset mechanisms are usually triggered passively when users approach a safety threshold. They prevent immediate physical collisions, but they do not directly evaluate recovery quality after a local conflict. This paper designs a proactive reset arbitration method for local conflict recovery. The method treats bilateral user resets, short-distance consecutive resets, and low post-conflict recovery distance as high-cost recovery events. It enters a proactive intervention window when a user pair still has positive recovery margin but the margin keeps decreasing. A local max-min arbitration rule then selects one user for unilateral reset. The reset direction follows the existing policy, so the study isolates the effect of trigger timing and user selection. Experiments with 3, 4, 5, and 6 users under APF_OSP show that Active Reset does not increase total reset rows and improves the median D_recover for all user counts. For 4, 5, and 6 users, short-distance recovery risks decrease significantly. The results suggest that proactive reset can replace part of the high-cost passive recovery process rather than simply adding extra interruptions.")
    add_para(doc, "Keywords: redirected walking; multi-user virtual reality; proactive reset; local conflict recovery; APF_OSP; recovery distance")


def add_background(doc):
    add_heading(doc, "1 研究背景与意义", 1)
    add_para(doc, "虚拟现实系统希望用户在虚拟环境中自然移动，但真实可用物理空间通常有限。重定向行走通过调整真实运动和虚拟运动之间的映射，使用户在有限空间中感知到更大的虚拟空间 [1,2]。与传送、原地行走和跑步机等方式相比，重定向行走更接近真实步行，因此常用于强调沉浸感和空间存在感的虚拟现实应用。")
    add_para(doc, "多人重定向行走进一步增加了控制难度。多个用户共享同一物理空间时，系统不仅要防止用户撞墙，还要防止用户之间发生真实碰撞。用户间冲突不是简单的距离过近问题。若两名用户正在自然分离，短距离接近可能不需要干预；若两名用户仍在相向运动，即使尚未到达安全阈值，也可能很快进入危险状态。")
    add_para(doc, "现有被动用户间重置主要解决即时安全问题。当距离过近或预测即将碰撞时，系统通过重置调整用户方向。该机制必要但不充分，因为它没有显式评价一次重置之后是否能稳定恢复。如果处理过晚，一次局部冲突可能引发两名用户都被重置，或某个用户刚完成重置后很快再次重置。本文的意义在于把这类现象定义为局部冲突恢复质量问题，并尝试用主动重置仲裁降低高代价恢复事件。")


def add_related(doc):
    add_heading(doc, "2 国内外研究现状", 1)
    add_heading(doc, "2.1 重定向行走与重置机制", 2)
    add_para(doc, "Razzaque 等提出的 redirected walking 奠定了 RDW 的基本框架 [1]。Steinicke 等测量了平移、旋转和曲率增益的感知阈值，为不可察觉重定向提供了基础 [2]。当隐式增益不足以保证安全时，系统需要显式重置。Williams 等比较了 Freeze-Backup、Freeze-Turn 和 2:1-Turn 等重置方式 [3]。Hodgson 等在受限虚拟环境中比较多种 RDW 算法，使用 reset count 和 distance between resets 等指标评价性能 [4]。这些工作说明，重置数量和两次重置之间的距离是 RDW 评价中的核心指标。")
    add_heading(doc, "2.2 多人 RDW 与用户间避碰", 2)
    add_para(doc, "多人 RDW 的核心问题是共享物理空间中的协同控制。Bachmann 等将人工势场用于多人 RDW，把障碍物和其他用户视为排斥源，从而减少用户碰撞和重置 [5]。Thomas 和 Rosenberg 提出通用的 APF reactive RDW 算法，为 APF 类控制器提供了更一般的形式 [6]。Dong 等的 FREE-RDW 进一步利用非前向步态的感知阈值处理多人碰撞规避 [11]。这些方法证明了多人 RDW 可以通过势场、运动模式和用户协同降低碰撞风险。")
    add_heading(doc, "2.3 空间划分、预测与学习方法", 2)
    add_para(doc, "随着问题复杂度提高，研究者开始引入未来状态和空间规划。Jeon 等提出动态最优空间划分 OSP，为不同用户分配物理子空间 [7]。Xu 等提出 Optimal Pose Guided 方法，通过离散化位置和朝向，并计算 pose 的长期安全分数，引导用户走向更安全的未来 pose [12]。Hirt 等将预测思想引入多用户 APF RDW [8]。Lee 等用强化学习设计多用户 reset controller，以优化 reset 决策 [9]。这些研究关注全局空间利用、未来轨迹和方向策略，而本文关注更局部的问题：一次用户间冲突是否会在恢复阶段演化为双边重置或短距离连续重置。")


def add_method(doc):
    add_heading(doc, "3 模型方法", 1)
    add_heading(doc, "3.1 数据采集与预处理", 2)
    add_para(doc, "本文使用 3、4、5、6 用户 APF_OSP 仿真实验处理结果。每个用户数下包含 Baseline 和 Active Reset 两种条件。分析脚本从实验日志中生成 normalized_reset_rows.csv 和 event_recovery_outcomes.csv 等派生文件。前者记录每次 reset 的类型，后者记录用户相关事件之后到下一次相关 reset/event 的恢复行走距离。")
    add_para(doc, "预处理主要包括三步。第一，统一 reset 类型，将 wall reset、ordinary USER_RESET 和 PROACTIVE_USER_RESET 分开统计。第二，按用户数和实验条件聚合 total reset rows、USER+PRO reset rows 和 mean distance between resets。第三，从用户相关事件中提取恢复距离，并计算 median、四分位数、D_recover <1m/<2m/<3m 的比例。episode-level 统计以 episode 为独立样本，每组 100 个 episode。")
    add_heading(doc, "3.2 问题定义", 2)
    add_para(doc, "本文将局部用户间冲突定义为：在一段局部时间窗口内，两个用户的物理距离、相对运动趋势或预测轨迹显示其可能进入安全风险区域，并可能触发用户间重置的 pair-level 状态。高代价恢复包括三种可观测形式：双边用户重置、短距离连续重置和低 D_recover。")
    add_heading(doc, "3.3 主动重置触发模型", 2)
    add_para(doc, "设 best.Margin 表示预测窗口内最优左右转组合下的最小安全余量。若 best.Margin < 0，说明即使局部最优转向也会越过安全边界，这更接近预测版被动安全处理。本文希望在更早阶段干预，因此定义恢复余量趋势窗口：0 < best.Margin <= clamp(closingSpeed * 0.30s + 0.05m, 0.15m, 0.50m)，且最近 5 帧中至少 3 帧 margin 下降。该窗口表示冲突尚未不可恢复，但正在逼近高代价恢复状态。")
    add_heading(doc, "3.4 单侧用户仲裁模型", 2)
    add_para(doc, "进入主动窗口后，系统分别估计重置用户 A 或用户 B 后的局部恢复效果。本文采用局部 max-min 思路，不选择“当前最危险”的用户，而选择能提高当前 pair 较差恢复状态的一侧用户。若被选方案的 worst remaining distance 未比保持当前状态至少提高 0.20m，则拒绝该候选。这样可以减少误触发，避免主动重置成为额外打断。")
    add_heading(doc, "3.5 算法流程", 2)
    for item in [
        "构建相邻用户 pair，并过滤静止、非相向或单侧运动导致的候选。",
        "计算 best.Margin、closing speed 和最近若干帧的 margin 下降趋势。",
        "若 pair 仍有正恢复余量且进入主动窗口，则创建主动重置触发事件。",
        "分别估计重置两侧用户后的 worst remaining distance，并用 max-min 规则选择用户。",
        "检查最小预期收益、冷却时间和安全约束；若通过，则生成主动重置意图。"
    ]:
        add_bullet(doc, item)


def add_results(doc):
    add_heading(doc, "4 实验结果与分析", 1)
    add_heading(doc, "4.1 总 reset rows 没有增加", 2)
    add_para(doc, "首先检查 Active Reset 是否只是额外插入更多 reset。结果显示，3、4、5、6 用户下，Active Reset 的 total reset rows 均低于 Baseline，分别为 7473 -> 7347、11758 -> 11341、17318 -> 16384、23870 -> 22966。ordinary USER_RESET 也明显下降，说明主动重置更像是替代了一部分被动用户间 reset，而不是简单增加额外打断。")
    add_table(doc, ["用户数", "条件", "Total rows", "Wall", "Ordinary USER", "Proactive", "USER+PRO", "Mean distance"],
              [["3", "Baseline", "7473", "6154", "1319", "0", "1319", "8.005 m"],
               ["3", "Active", "7347", "6154", "878", "315", "1193", "8.168 m"],
               ["4", "Baseline", "11758", "8612", "3146", "0", "3146", "6.871 m"],
               ["4", "Active", "11341", "8514", "2046", "781", "2827", "7.126 m"],
               ["5", "Baseline", "17318", "11271", "6047", "0", "6047", "5.911 m"],
               ["5", "Active", "16384", "11175", "3774", "1435", "5209", "6.222 m"],
               ["6", "Baseline", "23870", "14122", "9748", "0", "9748", "5.187 m"],
               ["6", "Active", "22966", "13893", "6699", "2374", "9073", "5.384 m"]],
              [1.2, 2.1, 2.0, 1.6, 2.1, 1.8, 1.8, 2.2])
    add_figure(doc, RESULT_DIR + r"\fig0_total_reset_rows_composition.png", "图 1  不同用户数下 reset rows 组成。Active Reset 没有增加总体 reset rows。")
    add_heading(doc, "4.2 恢复行走距离增加", 2)
    add_para(doc, "Active Reset 在所有用户数下均提高用户相关事件后的 median recovery distance。3 用户下从 3.262m 提高到 3.614m；4 用户下从 2.308m 提高到 3.329m；5 用户下从 2.120m 提高到 2.789m；6 用户下从 1.676m 提高到 2.199m。")
    add_para(doc, "以 2m 为阈值，短距离恢复比例在所有用户数下下降。Baseline 与 Active Reset 的 <2m 比例分别为：3 用户 32.22% -> 28.61%，4 用户 45.15% -> 34.64%，5 用户 48.18% -> 39.69%，6 用户 55.56% -> 46.92%。")
    add_table(doc, ["用户数", "Baseline median", "Active median", "Baseline <2m", "Active <2m"],
              [["3", "3.262 m", "3.614 m", "32.22%", "28.61%"],
               ["4", "2.308 m", "3.329 m", "45.15%", "34.64%"],
               ["5", "2.120 m", "2.789 m", "48.18%", "39.69%"],
               ["6", "1.676 m", "2.199 m", "55.56%", "46.92%"]],
              [2.0, 3.1, 3.1, 3.1, 3.1])
    add_figure(doc, RESULT_DIR + r"\fig7_user_collision_recovery_distance_survival.png", "图 2  用户相关事件后的恢复距离 survival curve。Active Reset 曲线整体高于 Baseline。")
    add_heading(doc, "4.3 episode-level 显著性分析", 2)
    add_para(doc, "episode-level permutation test 表明，D_recover median 在所有用户数下显著提升。提升量分别为 0.615m、0.903m、0.690m 和 0.568m。短距离风险指标在 4、5、6 用户下稳定显著下降；3 用户下两个 <2m share 指标下降方向一致，但未达到显著。")
    add_table(doc, ["用户数", "D_recover median", "Short <2m share", "Next reset <2m share"],
              [["3", "+0.615 m, p=0.0050", "-2.86 pp, p=0.2976", "-2.92 pp, p=0.2861"],
               ["4", "+0.903 m, p=0.0001", "-11.88 pp, p=0.0001", "-10.61 pp, p=0.0001"],
               ["5", "+0.690 m, p=0.0001", "-9.22 pp, p=0.0001", "-8.93 pp, p=0.0001"],
               ["6", "+0.568 m, p=0.0001", "-10.32 pp, p=0.0001", "-8.89 pp, p=0.0001"]],
              [1.6, 4.4, 4.4, 4.4])
    add_figure(doc, RESULT_DIR + r"\fig6_episode_violin_significance.png", "图 3  episode-level 分布和显著性。4/5/6 用户下短距离风险下降更稳定。")
    add_heading(doc, "4.4 结果解释", 2)
    add_para(doc, "这些结果支持本文的核心假设：主动重置的价值不是增加 reset 数量，而是把一部分高代价被动恢复转化为更可控的单侧干预。轻载 3 用户场景下，主动重置仍提升 D_recover median，但短距离比例下降不显著，说明该方法在冲突压力较低时可替代空间有限。4、5、6 用户下效果更稳定，说明方法更适合中高密度多人场景。")


def add_conclusion(doc):
    add_heading(doc, "5 总结与展望", 1)
    add_para(doc, "本文围绕多人 RDW 中的局部用户间冲突恢复问题，设计了一种主动重置仲裁方法。方法将高代价恢复定义为双边用户重置、短距离连续重置和低 D_recover，并在恢复余量仍为正但持续下降时触发主动干预。随后，系统通过局部 max-min 仲裁选择单侧用户重置，以改善当前 pair 的较差恢复状态。")
    add_para(doc, "实验结果表明，Active Reset 没有增加总 reset rows，并能提高用户相关事件后的恢复距离。4、5、6 用户下短距离恢复风险显著下降，说明主动重置在中高冲突密度场景中更稳定。该结果支持本文判断：主动重置可以替代部分高代价被动恢复，而不是简单插入额外打断。")
    add_para(doc, "后续工作可以从三个方向展开。第一，在更多 RDW 控制器和不同物理空间布局下验证泛化能力。第二，引入真实用户实验，评价主观舒适度、沉浸感和任务表现。第三，在保持本文触发与仲裁框架的基础上，再研究是否需要更复杂的重置方向优化。")


def add_references(doc):
    add_heading(doc, "参考文献", 1)
    refs = [
        "[1] Razzaque, S., Kohn, Z., Whitton, M. C. Redirected walking. In: Eurographics 2001 Short Presentations, 2001.",
        "[2] Steinicke, F., Bruder, G., Jerald, J., Frenz, H., Lappe, M. Estimation of detection thresholds for redirected walking techniques. IEEE Transactions on Visualization and Computer Graphics, 16(1), 17-27, 2010.",
        "[3] Williams, B., Narasimham, G., Rump, B., McNamara, T. P., Carr, T. H., Rieser, J. J., Bodenheimer, B. Exploring large virtual environments with an HMD when physical space is limited. In: APGV 2007, pp. 41-48, 2007.",
        "[4] Hodgson, E., Bachmann, E. R., Thrash, T. Performance of redirected walking algorithms in a constrained virtual world. IEEE Transactions on Visualization and Computer Graphics, 20(4), 579-587, 2014.",
        "[5] Bachmann, E. R., Hodgson, E., Hoffbauer, C., Messinger, J. Multi-user redirected walking and resetting using artificial potential fields. IEEE Transactions on Visualization and Computer Graphics, 25(5), 2022-2031, 2019.",
        "[6] Thomas, J., Rosenberg, E. S. A general reactive algorithm for redirected walking using artificial potential functions. In: IEEE VR 2019, pp. 56-62, 2019.",
        "[7] Jeon, S. B., Kwon, S. U., Hwang, J. Y., Cho, Y. H., Kim, H., Park, J., Lee, I. K. Dynamic optimal space partitioning for redirected walking in multi-user environment. ACM Transactions on Graphics, 41(4), Article 90, 2022.",
        "[8] Hirt, C., Isaak, N., Holz, C., Kunz, A. Predictive multiuser redirected walking using artificial potential fields. Frontiers in Virtual Reality, 5, Article 1259429, 2024.",
        "[9] Lee, H. J., Jeon, S.-B., Cho, Y.-H., Lee, I.-K. Multi-user reset controller for redirected walking using reinforcement learning. arXiv:2306.11433, 2023.",
        "[10] Liu, J.-H., Ren, Y.-F., Gan, Q. W., Huang, K., Chen, F. X. Y., Luo, E.-X., Tang, K. Y., Fu, Y.-Y., Fan, C.-W., Zhang, F.-L., Zhang, S.-H. A survey on redirected walking in virtual reality. IEEE Transactions on Visualization and Computer Graphics, 2024.",
        "[11] Dong, T., Gao, T., Dong, Y., Wang, L., Hu, K., Fan, J. FREE-RDW: A multi-user redirected walking method for supporting non-forward steps. IEEE Transactions on Visualization and Computer Graphics, 2024.",
        "[12] Xu, S.-Z., Lv, T., He, G., Chen, C.-H., Zhang, F.-L., Zhang, S.-H. Optimal pose guided redirected walking. ACM Transactions on Graphics, 41(4), Article 89, 2022.",
    ]
    for ref in refs:
        add_para(doc, ref)


def main():
    doc = Document()
    style_document(doc)
    add_front(doc)
    add_background(doc)
    add_related(doc)
    add_method(doc)
    add_results(doc)
    add_conclusion(doc)
    add_references(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
